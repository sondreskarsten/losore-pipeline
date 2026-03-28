import os
import re
import csv
import json
import gzip
import time
import requests
from datetime import datetime, timezone
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from google.cloud import storage


BUCKET_NAME = os.environ.get("GCS_BUCKET", "sondre_brreg_data")
GCS_PREFIX = os.environ.get("GCS_PREFIX", "losore/agder")

KOMMUNENUMMER = os.environ.get("KOMMUNENUMMER", ",".join([
    "4201", "4202", "4203", "4204", "4205", "4206", "4207",
    "4211", "4212", "4213", "4214", "4215", "4216", "4217",
    "4218", "4219", "4220", "4221", "4222", "4223", "4224",
    "4225", "4226", "4227", "4228",
])).split(",")

SEARCH_TERMS = os.environ.get("SEARCH_TERMS", "SPAREBANKEN NORGE,SPAREBANKEN SØR").split(",")

RELEVANT_CATEGORIES = [
    "Pant i driftstilbehør", "Pant i varelager", "Pant i landbruksredskaper",
    "Pant i fiskeredskaper", "Pant i fordringer (factoring)",
    "Pant i motorvogner, anleggsmaskiner og jernbanemateriell",
]

DELAY = float(os.environ.get("SCRAPE_DELAY", "0.05"))
SAVE_EVERY = int(os.environ.get("SAVE_EVERY", "100"))
ORGNR_MIN = os.environ.get("ORGNR_MIN", "")
ORGNR_MAX = os.environ.get("ORGNR_MAX", "")

GZ_FILE = "/tmp/enhetsregisteret_alle.csv.gz"
RAW_FILE = "/tmp/raw_responses.jsonl"
PARSED_FILE = "/tmp/parsed_data.json"


# ═══════════════════════════════════════════════════════════════
# GCS helpers
# ═══════════════════════════════════════════════════════════════

def gcs_client():
    """Return a google.cloud.storage.Client instance."""
    return storage.Client()


def gcs_upload(local_path, gcs_path):
    """Upload a local file to GCS.

    Parameters
    ----------
    local_path : str
        Local filesystem path to upload.
    gcs_path : str
        Full GCS object path (without ``gs://bucket/`` prefix).
    """
    client = gcs_client()
    bucket = client.bucket(BUCKET_NAME)
    blob = bucket.blob(gcs_path)
    blob.upload_from_filename(local_path)


def gcs_download(gcs_path, local_path):
    """Download a GCS object to a local file.

    Parameters
    ----------
    gcs_path : str
        GCS object path relative to ``BUCKET_NAME``.
    local_path : str
        Local destination path.

    Returns
    -------
    bool
        ``True`` if the blob existed and was downloaded, ``False`` otherwise.
    """
    client = gcs_client()
    bucket = client.bucket(BUCKET_NAME)
    blob = bucket.blob(gcs_path)
    if blob.exists():
        blob.download_to_filename(local_path)
        return True
    return False


def gcs_upload_json(data, gcs_path):
    """Serialise a Python object to JSON and upload to GCS.

    Parameters
    ----------
    data : dict or list
        JSON-serialisable data.
    gcs_path : str
        GCS object path.
    """
    client = gcs_client()
    bucket = client.bucket(BUCKET_NAME)
    blob = bucket.blob(gcs_path)
    blob.upload_from_string(json.dumps(data, ensure_ascii=False), content_type="application/json")


def update_status(phase, progress=None, detail=None, error=None):
    """Write pipeline status to ``{GCS_PREFIX}/status.json`` on GCS.

    Called at every phase transition. Includes pipeline configuration
    (bucket, prefix, kommunenummer, search terms) and timestamp.

    Parameters
    ----------
    phase : str
        Current phase (``"starting"``, ``"download"``, ``"collect"``,
        ``"parse"``, ``"export"``, ``"done"``, ``"error"``).
    progress : str or None
        Numeric progress indicator (e.g. ``"500/1200"``).
    detail : str or None
        Human-readable detail message.
    error : str or None
        Error message if phase is ``"error"``.
    """
    status = {
        "phase": phase,
        "progress": progress,
        "detail": detail,
        "error": error,
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "config": {
            "bucket": BUCKET_NAME,
            "prefix": GCS_PREFIX,
            "kommunenummer": KOMMUNENUMMER,
            "search_terms": SEARCH_TERMS,
        },
    }
    gcs_upload_json(status, f"{GCS_PREFIX}/status.json")
    print(f"[{phase}] {progress or ''} {detail or ''}", flush=True)


# ═══════════════════════════════════════════════════════════════
# Stage 0: Download enhetsregisteret
# ═══════════════════════════════════════════════════════════════

def download_enhetsregisteret():
    """Download the full enhetsregisteret CSV for kommune filtering.

    Three-level cache: local file → GCS cache → brreg API download.
    The ~152 MB gzipped CSV is used in stage 1 to filter orgnrs by
    ``KOMMUNENUMMER`` and in stage 3 to enrich output with entity
    metadata (name, legal form, NACE code, employees).

    Downloads from ``https://data.brreg.no/enhetsregisteret/api/enheter/lastned/csv``.
    """
    if os.path.exists(GZ_FILE):
        update_status("download", detail="Using cached local file")
        return
    if gcs_download(f"{GCS_PREFIX}/enhetsregisteret_alle.csv.gz", GZ_FILE):
        update_status("download", detail="Restored from GCS cache")
        return
    url = "https://data.brreg.no/enhetsregisteret/api/enheter/lastned/csv"
    update_status("download", detail=f"Downloading from {url}")
    resp = requests.get(url, stream=True, timeout=300)
    resp.raise_for_status()
    downloaded = 0
    with open(GZ_FILE, "wb") as f:
        for chunk in resp.iter_content(chunk_size=1024 * 1024):
            f.write(chunk)
            downloaded += len(chunk)
    gcs_upload(GZ_FILE, f"{GCS_PREFIX}/enhetsregisteret_alle.csv.gz")
    update_status("download", detail=f"Downloaded {downloaded / 1e6:.1f} MB, cached to GCS")


# ═══════════════════════════════════════════════════════════════
# Stage 1: COLLECT — JSONL streaming, O(batch_size) memory
#
# Primary: RSC endpoint (Rsc: 1 header) → text/x-component
#   - Returns flight stream directly, no HTML wrapper
#   - No JS string unescape needed — lines are plain text
#   - ~18% the size of full HTML response
#   - Eliminates the escape-bug class entirely
#
# Fallback: HTML parsing (self.__next_f.push extraction)
#   - Used if RSC endpoint returns non-200 or wrong content-type
#   - Requires regex extraction + json.loads JS unescape
#
# Record schema:
#   orgnr            str
#   url              str
#   collected_at     ISO timestamp
#   http_status      int | null
#   rsc_payload      str (flight stream line or decoded push payload)
#   rsc_payload_raw  str (undecoded push payload, fallback only)
#   method           "rsc" | "html" | null
#   error            str | null
# ═══════════════════════════════════════════════════════════════

def collect_one(orgnr, session):
    """Collect løsøreregisteret data for one orgnr.

    Two-strategy collection:

    1. **RSC endpoint** (primary): sends ``Rsc: 1`` header to get
       React Server Component flight stream (``text/x-component``).
       Searches for the line containing ``"rettsstiftelser":[``.
       ~18% the size of full HTML, no JS unescape needed.

    2. **HTML fallback**: if RSC fails, fetches the full page and
       extracts ``self.__next_f.push([1,"..."])`` payloads via regex.
       Requires ``json.loads`` to unescape JS string encoding.

Parameters
    ----------
    orgnr : str
        9-digit Norwegian organisation number.
    session : requests.Session
        HTTP session with connection pooling and user-agent header.

    Returns
    -------
    dict
        Record with keys: ``orgnr``, ``url``, ``collected_at``,
        ``http_status``, ``rsc_payload`` (str or None),
        ``rsc_payload_raw`` (str or None), ``method`` (``"rsc"`` or
        ``"html"`` or None), ``error`` (str or None).
    """
    url = f"https://rettsstiftelser.brreg.no/nb/oppslag/virksomhet/{orgnr}"
    record = {
        "orgnr": orgnr,
        "url": url,
        "collected_at": datetime.now(timezone.utc).isoformat(),
        "http_status": None,
        "rsc_payload": None,
        "rsc_payload_raw": None,
        "method": None,
        "error": None,
    }

    try:
        resp = session.get(url, headers={"Rsc": "1"}, timeout=30)
        record["http_status"] = resp.status_code
        resp.raise_for_status()
        if resp.headers.get("content-type", "").startswith("text/x-component"):
            record["method"] = "rsc"
            for line in resp.content.decode("utf-8").split('\n'):
                if '"rettsstiftelser":[' in line:
                    record["rsc_payload"] = line
                    break
            return record
    except (requests.exceptions.RequestException,):
        pass

    record["method"] = "html"
    try:
        resp = session.get(url, timeout=30)
        record["http_status"] = resp.status_code
        resp.raise_for_status()
    except requests.exceptions.Timeout:
        record["error"] = "timeout"
        return record
    except requests.exceptions.ConnectionError as e:
        record["error"] = f"connection: {str(e)[:200]}"
        return record
    except requests.exceptions.HTTPError:
        record["error"] = f"http_{resp.status_code}"
        return record

    pushes = re.findall(
        r'self\.__next_f\.push\(\[1,"(.*?)"\]\)', resp.text, re.DOTALL
    )

    for payload in pushes:
        if 'rettsstiftelser' not in payload:
            continue
        try:
            decoded = json.loads('"' + payload + '"')
            record["rsc_payload"] = decoded
        except (json.JSONDecodeError, ValueError):
            record["rsc_payload_raw"] = payload
        break

    return record


def scan_existing_orgnr(jsonl_path):
    """Scan a JSONL file and return orgnrs that have been collected.

    Reads line by line, considers an orgnr "collected" if it has
    a non-null ``rsc_payload``, ``rsc_payload_raw``, or ``error``.
    Used for resume-on-restart: skip already-collected orgnrs.

    Parameters
    ----------
    jsonl_path : str
        Path to the raw_responses.jsonl file.

    Returns
    -------
    set[str]
        Orgnrs with existing collection records.
    """
    already = set()
    if not os.path.exists(jsonl_path):
        return already
    with open(jsonl_path, "r") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
                orgnr = rec["orgnr"]
                if rec.get("rsc_payload") is not None or rec.get("rsc_payload_raw") is not None:
                    already.add(orgnr)
                elif rec.get("error"):
                    already.add(orgnr)
            except (json.JSONDecodeError, KeyError):
                continue
    return already


def collect_all():
    """Stage 1: collect løsøreregisteret data for all target orgnrs.

    Filters the enhetsregisteret CSV by ``KOMMUNENUMMER``, optionally
    by ``ORGNR_MIN``/``ORGNR_MAX``, then scrapes each orgnr via
    ``collect_one()``.  Results stream to ``raw_responses.jsonl``
    in append-only JSONL format.  Checkpoints to GCS every
    ``SAVE_EVERY`` records for crash recovery.

    Resume: reads existing JSONL on startup, skips already-collected
    orgnrs via ``scan_existing_orgnr()``.
    """
    checkpoint_gcs = f"{GCS_PREFIX}/raw_responses.jsonl"
    if not os.path.exists(RAW_FILE):
        gcs_download(checkpoint_gcs, RAW_FILE)

    already = scan_existing_orgnr(RAW_FILE)
    update_status("collect", detail=f"Restored checkpoint: {len(already)} orgnr")

    target_orgnr = set()
    with gzip.open(GZ_FILE, "rt", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row.get("forretningsadresse.kommunenummer") in KOMMUNENUMMER:
                orgnr = row["organisasjonsnummer"]
                if ORGNR_MIN and orgnr < ORGNR_MIN:
                    continue
                if ORGNR_MAX and orgnr >= ORGNR_MAX:
                    continue
                target_orgnr.add(orgnr)

    remaining = sorted(target_orgnr - already)
    total_target = len(target_orgnr)
    update_status("collect", progress=f"{len(already)}/{total_target}",
                  detail=f"Remaining: {len(remaining)}")

    session = requests.Session()
    session.headers.update({
        "User-Agent": "SparebankenNorge-LosoreAnalyse/1.0 (+https://sparebanken.no)",
        "Accept-Encoding": "gzip",
    })

    batch = []
    for i, orgnr in enumerate(remaining, 1):
        record = collect_one(orgnr, session)
        batch.append(json.dumps(record, ensure_ascii=False))

        if i % SAVE_EVERY == 0:
            with open(RAW_FILE, "a") as f:
                f.write("\n".join(batch) + "\n")
            batch.clear()
            gcs_upload(RAW_FILE, checkpoint_gcs)
            done = len(already) + i
            update_status("collect", progress=f"{done}/{total_target}",
                          detail=f"Last: {orgnr}")

        time.sleep(DELAY)

    if batch:
        with open(RAW_FILE, "a") as f:
            f.write("\n".join(batch) + "\n")
        batch.clear()
        gcs_upload(RAW_FILE, checkpoint_gcs)

    update_status("collect", progress=f"{total_target}/{total_target}", detail="Complete")


# ═══════════════════════════════════════════════════════════════
# Stage 2: PARSE — stream JSONL, build parsed output
#
# Reads raw_responses.jsonl line by line. Builds parsed_data.json
# in memory (parsed entries are much smaller than raw payloads).
# ═══════════════════════════════════════════════════════════════

def extract_rettsstiftelser(rsc_payload):
    """Extract the rettsstiftelser array from an RSC payload string.

    Handles encoding recovery: tries ``latin-1 → utf-8`` re-encoding
    for payloads that were double-encoded during collection.  Parses
    the ``"data":{...}`` JSON object via brace-depth counting (the
    payload is not valid standalone JSON — it's a React flight stream
    line with surrounding metadata).

    Parameters
    ----------
    rsc_payload : str
        Raw RSC flight stream line or decoded HTML push payload
        containing ``"data":{"rettsstiftelser":[...]}``.

    Returns
    -------
    list[dict] or None
        List of rettsstiftelse dicts, or ``None`` if no data block
        found.  Each dict has keys: ``dokumentnummer``,
        ``typeBeskrivelse``, ``statusBeskrivelse``,
        ``innkomsttidspunkt``, ``roller``, ``formuesgoder``, ``krav``,
        ``paategninger``, ``prioritetsvikelser``, ``konkurs``.
    """
    try:
        rsc_payload = rsc_payload.encode("latin-1").decode("utf-8")
    except (UnicodeDecodeError, UnicodeEncodeError):
        pass
    match = re.search(r'"data":\s*(\{)', rsc_payload)
    if not match:
        return None
    raw = rsc_payload[match.start(1):]
    depth = 0
    for i, ch in enumerate(raw):
        if ch == '{':
            depth += 1
        elif ch == '}':
            depth -= 1
        if depth == 0:
            return json.loads(raw[:i + 1]).get("rettsstiftelser", [])
    return None


def parse_entry(rs):
    """Parse a single rettsstiftelse dict into a flat entry.

    Extracts: title (``typeBeskrivelse``), document number, timestamp,
    status, and for each rolle: panthaver/pantsetter/eier with name
    and orgnr.  Formuesgoder are concatenated into a single string
    with ``" | "`` separator.  Krav beløp is extracted as integer.
    Påtegninger are concatenated.

    Parameters
    ----------
    rs : dict
        One rettsstiftelse from ``extract_rettsstiftelser()``.

    Returns
    -------
    dict or None
        Flat entry dict, or ``None`` if input is not a dict.
    """
    if not isinstance(rs, dict):
        return None

    entry = {
        "title": rs.get("typeBeskrivelse"),
        "document_number": rs.get("dokumentnummer"),
        "innkomsttidspunkt": rs.get("innkomsttidspunkt"),
        "rettsstiftelsen_er": rs.get("statusBeskrivelse"),
    }

    for rolle in rs.get("roller", []):
        if not isinstance(rolle, dict):
            continue
        ri = rolle.get("rolleinnehaver", {})
        if not isinstance(ri, dict):
            continue
        key = rolle.get("rolletype", "")
        if key == "rolletype.panthaver":
            entry["panthaver"] = ri.get("navn", "")
            entry["panthaver_orgnr"] = ri.get("organisasjonsnummer", "")
        elif key == "rolletype.pantsetter":
            entry["pantsetter"] = ri.get("navn", "")
            entry["pantsetter_orgnr"] = ri.get("organisasjonsnummer", "")
        elif key == "rolletype.eier":
            entry["eier"] = ri.get("navn", "")
            entry["eier_orgnr"] = ri.get("organisasjonsnummer", "")

    fgs = []
    for fg in rs.get("formuesgoder", []):
        if not isinstance(fg, dict):
            continue
        desc = fg.get("typeBeskrivelse", "")
        avgrensing = fg.get("avgrensingTingsinnbegrepBeskrivelse", "")
        if avgrensing:
            desc += f" : {avgrensing}"
        ident = fg.get("identifisering", "")
        if ident:
            desc += f" {ident}"
        regnr = fg.get("registreringsnummer", "")
        if regnr:
            desc += f" {regnr}"
        fgs.append(desc)
    entry["formuesgode"] = " | ".join(fgs) if fgs else ""

    krav = rs.get("krav")
    if isinstance(krav, dict):
        belop_list = krav.get("belop", [])
        if belop_list and isinstance(belop_list, list) and len(belop_list) > 0 and isinstance(belop_list[0], dict):
            entry["beløp"] = int(belop_list[0].get("belop", 0))
        else:
            entry["beløp"] = None
        entry["krav"] = krav.get("kravFordringerBeskrivelse", "") or krav.get("beskrivelse", "")
    else:
        entry["beløp"] = None
        entry["krav"] = ""

    paategninger = rs.get("paategninger", [])
    if paategninger and isinstance(paategninger, list):
        pts = [pt.get("paategning", "") for pt in paategninger if isinstance(pt, dict)]
        if pts:
            entry["påtegning"] = " | ".join(pts)

    return entry


def parse_record(record):
    """Parse a raw collection record into structured entries.

    Handles both current (``rsc_payload``) and legacy
    (``rsc_payloads`` list) record formats.  Tries decoded payloads
    first, falls back to ``rsc_payload_raw`` with JS unescape.

    Parameters
    ----------
    record : dict
        One line from ``raw_responses.jsonl``.

    Returns
    -------
    dict
        Parsed record with keys: ``orgnr``, ``entries`` (list),
        ``entries_count``, ``raw_rettsstiftelser`` (list),
        ``collect_error``, ``parse_error``.
    """
    result = {
        "orgnr": record["orgnr"],
        "entries": [],
        "entries_count": 0,
        "raw_rettsstiftelser": [],
        "url": record.get("url", ""),
        "collected_at": record.get("collected_at", ""),
        "collect_error": record.get("error"),
        "parse_error": None,
    }

    if record.get("error"):
        return result

    payload = record.get("rsc_payload")
    raw_payload = record.get("rsc_payload_raw")

    # Also handle old multi-payload format for backward compatibility
    payloads = record.get("rsc_payloads", [])
    raw_payloads = record.get("rsc_payloads_raw", [])

    sources = []
    if payload:
        sources.append(("decoded", payload))
    for p in payloads:
        sources.append(("decoded", p))
    if raw_payload:
        sources.append(("raw", raw_payload))
    for p in raw_payloads:
        sources.append(("raw", p))

    for source_type, pl in sources:
        try:
            if source_type == "raw":
                pl = json.loads('"' + pl + '"')
            rettsstiftelser = extract_rettsstiftelser(pl)
            if rettsstiftelser is None:
                continue
            result["raw_rettsstiftelser"] = rettsstiftelser
            for rs in rettsstiftelser:
                entry = parse_entry(rs)
                if entry:
                    result["entries"].append(entry)
            result["entries_count"] = len(result["entries"])
            break
        except Exception as e:
            result["parse_error"] = str(e)[:500]

    return result


def parse_all():
    """Stage 2: parse all raw responses into structured entries.

    Reads ``raw_responses.jsonl``, calls ``parse_record()`` per line,
    writes ``parsed_data.json`` (full dict keyed by orgnr).
    Tracks collect_error and parse_error counts separately.

    Returns
    -------
    dict[str, dict]
        Mapping from orgnr to parsed record.
    """
    update_status("parse", detail="Parsing raw responses")
    parsed = {}
    parse_errors = 0
    collect_errors = 0
    total = 0

    with open(RAW_FILE, "r") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                continue

            total += 1
            result = parse_record(record)
            parsed[result["orgnr"]] = result

            if result.get("collect_error"):
                collect_errors += 1
            if result.get("parse_error"):
                parse_errors += 1

    with open(PARSED_FILE, "w") as f:
        json.dump(parsed, f, ensure_ascii=False)
    gcs_upload(PARSED_FILE, f"{GCS_PREFIX}/parsed_data.json")

    with_entries = sum(1 for v in parsed.values() if v["entries_count"] > 0)
    update_status("parse", progress=f"{with_entries}/{total}",
                  detail=f"collect_errors={collect_errors} parse_errors={parse_errors}")
    return parsed


# ═══════════════════════════════════════════════════════════════
# Stage 3a: EXPORT ALL — structured flattened datasets per type
#
# Operates on raw_rettsstiftelser (full JSON). Data-driven:
# discovers all rolle types, formuesgode fields, krav fields
# from the data itself. Produces one sheet per entry type.
# ═══════════════════════════════════════════════════════════════

ROLLE_ADDRESS_FIELDS = [
    ("kommune", "ustrukturertadresse.kommune.kommunenavn"),
    ("kommunenr", "ustrukturertadresse.kommune.kommunenummer"),
    ("postnr", "ustrukturertadresse.poststed.postnummer"),
    ("poststed", "ustrukturertadresse.poststed.navn"),
    ("adresse", "ustrukturertadresse.adresse"),
]


def _deep_get(obj, dotpath):
    """Navigate a nested dict via dot-separated path.

    Parameters
    ----------
    obj : dict
        Root object.
    dotpath : str
        Dot-separated key path (e.g. ``"ustrukturertadresse.kommune.kommunenavn"``).

    Returns
    -------
    Any or None
    """
    for key in dotpath.split("."):
        if not isinstance(obj, dict):
            return None
        obj = obj.get(key)
    return obj


def flatten_rettsstiftelse(rs, orgnr):
    """Flatten a single rettsstiftelse into a wide row for Excel export.

    Extracts all fields including: per-rolle address details (kommune,
    postnr, adresse, international address), per-formuesgode details
    (type, avgrensing, regnr, VIN, eierandel, avtaletype),
    prioritetsvikelser, påtegninger, konkurs/tvangsavvikling flag.
    Timestamps parsed to datetime objects for Excel date formatting.

    Parameters
    ----------
    rs : dict
        One rettsstiftelse from the raw JSON.
    orgnr : str
        Organisation number.

    Returns
    -------
    dict or None
        Wide row dict with dynamic columns per rolle type and
        formuesgode index, or ``None`` if input is not a dict.
    """
    if not isinstance(rs, dict):
        return None

    row = {
        "orgnr": orgnr,
        "type": rs.get("typeBeskrivelse", ""),
        "dokumentnummer": rs.get("dokumentnummer", ""),
        "status": rs.get("statusBeskrivelse", ""),
    }

    ts = rs.get("innkomsttidspunkt", "")
    if ts:
        try:
            p = datetime.fromisoformat(ts)
            row["innkomst_dato"] = p.replace(tzinfo=None)
            if p.hour or p.minute:
                row["innkomst_kl"] = p.strftime("%H:%M")
        except ValueError:
            row["innkomst_dato"] = ts

    bts = rs.get("beslutningstidspunkt", "")
    if bts:
        try:
            p = datetime.fromisoformat(bts)
            row["beslutningsdato"] = p.replace(tzinfo=None)
        except ValueError:
            row["beslutningsdato"] = bts

    for rolle in rs.get("roller", []):
        if not isinstance(rolle, dict):
            continue
        rt = rolle.get("rolletype", "").replace("rolletype.", "")
        if not rt:
            continue
        ri = rolle.get("rolleinnehaver", {})
        if not isinstance(ri, dict):
            continue
        row[f"{rt}_navn"] = ri.get("navn", "")
        row[f"{rt}_orgnr"] = ri.get("organisasjonsnummer", "")
        row[f"{rt}_aktortype"] = ri.get("aktorType", "")
        for short, dotpath in ROLLE_ADDRESS_FIELDS:
            val = _deep_get(ri, dotpath)
            if isinstance(val, list):
                val = ", ".join(str(v) for v in val if v)
            if val:
                row[f"{rt}_{short}"] = val
        intl = ri.get("internasjonaladresse")
        if isinstance(intl, dict):
            row[f"{rt}_landkode"] = intl.get("landkode", "")
            row[f"{rt}_intl_adresse"] = intl.get("friAdressetekst", "") or intl.get("adressenavn", "")

    krav = rs.get("krav")
    if isinstance(krav, dict):
        belop_list = krav.get("belop", [])
        if belop_list and isinstance(belop_list, list) and len(belop_list) > 0 and isinstance(belop_list[0], dict):
            row["beløp"] = belop_list[0].get("belop")
            row["valuta"] = belop_list[0].get("valuta", "")
        row["krav_fordringer"] = krav.get("kravFordringerBeskrivelse", "")
        row["krav_salgspant"] = krav.get("kravSalgspantBeskrivelse", "")

    fgs = rs.get("formuesgoder", [])
    fg_descs = []
    for fi, fg in enumerate(fgs):
        if not isinstance(fg, dict):
            continue
        desc = fg.get("typeBeskrivelse", "")
        avg = fg.get("avgrensingTingsinnbegrepBeskrivelse", "")
        if avg:
            desc += f" : {avg}"
        ident = fg.get("beskrivelse", "")
        if ident:
            desc += f" {ident}"
        fg_descs.append(desc)

        prefix = f"fg{fi+1}_"
        row[f"{prefix}type"] = fg.get("typeBeskrivelse", "")
        row[f"{prefix}avgrensing"] = fg.get("avgrensingTingsinnbegrepBeskrivelse", "")
        row[f"{prefix}regnr"] = fg.get("registreringsnummerMotorvogn", "")
        hist = fg.get("historiskRegistreringsnummerMotorvogn", [])
        if isinstance(hist, list) and hist:
            row[f"{prefix}hist_regnr"] = ", ".join(str(h) for h in hist)
        row[f"{prefix}beskrivelse"] = fg.get("beskrivelse", "")
        row[f"{prefix}merke"] = fg.get("uregistrertMotorvognMerke", "")
        row[f"{prefix}aarsmodell"] = fg.get("uregistrertMotorvognAarsmodell", "")
        row[f"{prefix}vin"] = fg.get("uregistrertMotorvognIdentifikasjonsnummer", "")
        row[f"{prefix}identifiseringstype"] = fg.get("identifiseringstypeBeskrivelse", "")
        row[f"{prefix}identifikator"] = fg.get("identifikator", "")
        eierandel = fg.get("eierandel", {})
        if isinstance(eierandel, dict) and eierandel.get("teller") is not None:
            row[f"{prefix}eierandel"] = f"{eierandel['teller']}/{eierandel['nevner']}"
        row[f"{prefix}avtaletype_fordring"] = fg.get("avtaletypeFordringBeskrivelse", "")
        row[f"{prefix}org"] = fg.get("organisasjonsnummer", "")

    row["formuesgoder_antall"] = len(fg_descs)
    row["formuesgoder_beskrivelse"] = " | ".join(fg_descs) if fg_descs else ""

    pvs = rs.get("prioritetsvikelser", [])
    if pvs and isinstance(pvs, list):
        pv_docs = []
        for pv in pvs:
            if isinstance(pv, dict) and pv.get("dokumentnummer"):
                pv_docs.append(str(pv["dokumentnummer"]))
        if pv_docs:
            row["prioritetsvikelser_dok"] = ", ".join(pv_docs)

    pts = rs.get("paategninger", [])
    if pts and isinstance(pts, list):
        pt_texts = [pt.get("paategning", "") for pt in pts if isinstance(pt, dict)]
        if pt_texts:
            row["påtegninger"] = " | ".join(pt_texts)

    konkurs = rs.get("konkurs")
    if isinstance(konkurs, dict):
        row["er_tvangsavvikling"] = konkurs.get("erTvangsavviklingEllerTvangsopplosning", "")

    return row


def export_all(parsed):
    """Stage 3a: export all rettsstiftelser to a multi-sheet Excel file.

    One sheet per entry type (e.g. ``"Pant i driftstilbehør"``,
    ``"Utlegg"``), sorted by frequency.  Each sheet has entity
    metadata from enhetsregisteret joined by orgnr.  Column layout:
    entity fields → metadata → rolle details → krav → formuesgoder
    summary → formuesgoder detail → misc.  Empty columns auto-removed.

    Output: ``{GCS_PREFIX}/output/rettsstiftelser_all.xlsx``.

    Parameters
    ----------
    parsed : dict[str, dict]
        Output of ``parse_all()``.
    """
    update_status("export_all", detail="Building per-type datasets")

    by_type = defaultdict(list)
    for orgnr, d in parsed.items():
        for rs in d.get("raw_rettsstiftelser", []):
            row = flatten_rettsstiftelse(rs, orgnr)
            if row:
                by_type[row["type"]].append(row)

    ENHET_FIELDS = [
        "navn", "organisasjonsform.kode", "naeringskode1.kode", "naeringskode1.beskrivelse",
        "antallAnsatte", "forretningsadresse.kommune", "forretningsadresse.postnummer",
        "forretningsadresse.poststed", "stiftelsesdato", "sisteInnsendteAarsregnskap",
        "konkurs", "underAvvikling", "erIKonsern",
    ]
    all_orgnr = set()
    for rows in by_type.values():
        for r in rows:
            all_orgnr.add(r["orgnr"])

    enhet_lookup = {}
    with gzip.open(GZ_FILE, "rt", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            org = row.get("organisasjonsnummer")
            if org in all_orgnr:
                enhet_lookup[org] = {k: row.get(k, "") for k in ENHET_FIELDS}

    for rows in by_type.values():
        for r in rows:
            enhet = enhet_lookup.get(r["orgnr"], {})
            for k in ENHET_FIELDS:
                val = enhet.get(k, "")
                if k == "antallAnsatte" and val:
                    try:
                        val = int(val)
                    except ValueError:
                        pass
                r[k] = val

    enhet_cols = ["orgnr"] + ENHET_FIELDS
    meta_cols = ["type", "dokumentnummer", "status", "innkomst_dato", "innkomst_kl", "beslutningsdato"]
    krav_cols = ["beløp", "valuta", "krav_fordringer", "krav_salgspant"]
    fg_summary_cols = ["formuesgoder_antall", "formuesgoder_beskrivelse"]
    misc_cols = ["prioritetsvikelser_dok", "påtegninger", "er_tvangsavvikling"]

    fmt = {
        "beløp": '#,##0',
        "innkomst_dato": "DD.MM.YYYY",
        "beslutningsdato": "DD.MM.YYYY",
        "innkomst_kl": "@",
    }

    total_rows = 0
    wb = Workbook()
    wb.remove(wb.active)

    for entry_type in sorted(by_type.keys(), key=lambda t: -len(by_type[t])):
        rows = by_type[entry_type]

        ROLLE_SUFFIXES = ["_intl_adresse", "_landkode", "_aktortype", "_kommunenr",
                          "_kommune", "_poststed", "_postnr", "_adresse", "_orgnr", "_navn"]
        rolle_keys = set()
        fg_max = 0
        for r in rows:
            for k in r.keys():
                for sfx in ROLLE_SUFFIXES:
                    if k.endswith(sfx):
                        rolle_keys.add(k[:-len(sfx)])
                        break
            for i in range(1, 20):
                if f"fg{i}_type" in r and r[f"fg{i}_type"]:
                    fg_max = max(fg_max, i)

        rolle_types_ordered = sorted(rolle_keys, key=lambda rt: -sum(1 for r in rows if r.get(f"{rt}_navn")))
        rolle_cols = []
        for rt in rolle_types_ordered:
            rolle_cols.extend([
                f"{rt}_navn", f"{rt}_orgnr", f"{rt}_aktortype",
                f"{rt}_kommune", f"{rt}_kommunenr", f"{rt}_postnr", f"{rt}_poststed", f"{rt}_adresse",
                f"{rt}_landkode", f"{rt}_intl_adresse",
            ])

        fg_detail_cols = []
        for i in range(1, fg_max + 1):
            fg_detail_cols.extend([
                f"fg{i}_type", f"fg{i}_avgrensing", f"fg{i}_regnr", f"fg{i}_hist_regnr",
                f"fg{i}_beskrivelse", f"fg{i}_merke", f"fg{i}_aarsmodell", f"fg{i}_vin",
                f"fg{i}_identifiseringstype", f"fg{i}_identifikator", f"fg{i}_eierandel",
                f"fg{i}_avtaletype_fordring", f"fg{i}_org",
            ])

        all_cols = enhet_cols + meta_cols + rolle_cols + krav_cols + fg_summary_cols + fg_detail_cols + misc_cols
        present_cols = [c for c in all_cols if any(r.get(c) not in (None, "", 0) for r in rows)]

        sheet_name = entry_type[:31].replace("/", "-")
        ws = wb.create_sheet(title=sheet_name)

        hfont = Font(bold=True, size=10, name="Arial")
        hfill = PatternFill("solid", fgColor="D9E1F2")
        halign = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cfont = Font(size=10, name="Arial")

        for ci, h in enumerate(present_cols, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = hfont
            c.fill = hfill
            c.alignment = halign

        for ri, row in enumerate(rows, 2):
            for ci, h in enumerate(present_cols, 1):
                c = ws.cell(row=ri, column=ci, value=row.get(h))
                c.font = cfont

        for col_name, nf in fmt.items():
            if col_name in present_cols:
                ci = present_cols.index(col_name) + 1
                for ri in range(2, len(rows) + 2):
                    ws.cell(row=ri, column=ci).number_format = nf

        for ci, h in enumerate(present_cols, 1):
            sample_lens = [len(str(ws.cell(row=ri, column=ci).value or ""))
                           for ri in range(2, min(len(rows) + 2, 100))]
            max_len = max([len(h)] + sample_lens) if sample_lens else len(h)
            ws.column_dimensions[get_column_letter(ci)].width = min(max(10, max_len + 2), 40)

        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        total_rows += len(rows)

    local_xlsx = "/tmp/output_all_types.xlsx"
    wb.save(local_xlsx)
    gcs_xlsx = f"{GCS_PREFIX}/output/rettsstiftelser_all.xlsx"
    gcs_upload(local_xlsx, gcs_xlsx)
    update_status("export_all", progress=f"{total_rows} rows, {len(by_type)} types",
                  detail=f"gs://{BUCKET_NAME}/{gcs_xlsx}")


# ═══════════════════════════════════════════════════════════════
# Stage 3b: EXPORT FILTERED — Sparebanken Sør/Norge matches
# ═══════════════════════════════════════════════════════════════

def match_panthaver(entry):
    """Test whether an entry's panthaver/eier matches search terms.

    Parameters
    ----------
    entry : dict
        Parsed entry from ``parse_entry()``.

    Returns
    -------
    bool
    """
    ph = entry.get("panthaver", "") or entry.get("eier", "")
    return any(t.lower() in ph.lower() for t in SEARCH_TERMS)


def collect_rows(parsed):
    """Filter parsed data for entries matching SEARCH_TERMS and RELEVANT_CATEGORIES.

    Computes per-orgnr aggregates (total pant count and beløp for
    the matching bank) alongside individual entry rows.

    Parameters
    ----------
    parsed : dict[str, dict]
        Output of ``parse_all()``.

    Returns
    -------
    list[dict]
        Flat rows ready for Excel export.
    """
    belop_by_orgnr = defaultdict(int)
    count_by_orgnr = defaultdict(int)
    for orgnr, d in parsed.items():
        for e in d.get("entries", []):
            if not match_panthaver(e):
                continue
            if e.get("title", "") not in RELEVANT_CATEGORIES:
                continue
            b = e.get("beløp")
            if b:
                belop_by_orgnr[orgnr] += b
            count_by_orgnr[orgnr] += 1

    rows = []
    for orgnr, d in parsed.items():
        for e in d.get("entries", []):
            if not match_panthaver(e):
                continue
            if e.get("title", "") not in RELEVANT_CATEGORIES:
                continue
            ts = e.get("innkomsttidspunkt", "")
            dt = None
            kl = None
            if ts:
                try:
                    p = datetime.fromisoformat(ts)
                    dt = p.replace(tzinfo=None)
                    if p.hour or p.minute:
                        kl = p.strftime("%H:%M")
                except ValueError:
                    pass

            rows.append({
                "orgnr": orgnr,
                "title": e.get("title"),
                "panthaver": e.get("panthaver", ""),
                "panthaver_orgnr": e.get("panthaver_orgnr", ""),
                "pantsetter": e.get("pantsetter", ""),
                "pantsetter_orgnr": e.get("pantsetter_orgnr", ""),
                "beløp": e.get("beløp"),
                "document_number": e.get("document_number"),
                "innkomst_dato": dt,
                "innkomst_kl": kl,
                "rettsstiftelsen_er": e.get("rettsstiftelsen_er"),
                "formuesgode": e.get("formuesgode", ""),
                "krav": e.get("krav", ""),
                "total_pant_hos_bank": count_by_orgnr[orgnr],
                "total_beløp_hos_bank": belop_by_orgnr[orgnr],
            })
    return rows


def write_xlsx(rows, headers, output_file, fmt_map=None):
    """Write rows to a formatted Excel file.

    Applies header styling (bold, blue fill), number formatting per
    ``fmt_map``, auto-column-width, autofilter, frozen header row.

    Parameters
    ----------
    rows : list[dict]
        Data rows.
    headers : list[str]
        Column order.
    output_file : str
        Local path to write .xlsx.
    fmt_map : dict or None
        Column name → Excel number format (e.g. ``'#,##0" NOK"'``).

    Returns
    -------
    int
        Number of rows written.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Løsøre"
    hfont = Font(bold=True, size=11, name="Arial")
    hfill = PatternFill("solid", fgColor="D9E1F2")
    halign = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cfont = Font(size=11, name="Arial")
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = hfont
        c.fill = hfill
        c.alignment = halign
    for ri, row in enumerate(rows, 2):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=ri, column=ci, value=row.get(h))
            c.font = cfont
    if fmt_map:
        for col_name, nf in fmt_map.items():
            if col_name in headers:
                ci = headers.index(col_name) + 1
                for ri in range(2, len(rows) + 2):
                    ws.cell(row=ri, column=ci).number_format = nf
    for ci, h in enumerate(headers, 1):
        sample_lens = [len(str(ws.cell(row=ri, column=ci).value or ""))
                       for ri in range(2, min(len(rows) + 2, 100))]
        max_len = max([len(h)] + sample_lens) if sample_lens else len(h)
        ws.column_dimensions[get_column_letter(ci)].width = min(max(12, max_len + 3), 42)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    wb.save(output_file)
    return len(rows)


def export(parsed):
    """Stage 3b: export filtered Sparebanken entries with entity enrichment.

    Filters for ``SEARCH_TERMS`` × ``RELEVANT_CATEGORIES``, joins
    enhetsregisteret metadata (name, legal form, NACE, employees,
    address, bankruptcy status), writes formatted Excel to
    ``{GCS_PREFIX}/output/sparebanken_enhet.xlsx``.

    Parameters
    ----------
    parsed : dict[str, dict]
        Output of ``parse_all()``.
    """
    update_status("export", detail="Building xlsx")
    rows = collect_rows(parsed)
    fmt = {
        "beløp": '#,##0" NOK"',
        "total_beløp_hos_bank": '#,##0" NOK"',
        "innkomst_dato": "DD.MM.YYYY",
        "innkomst_kl": "@",
    }

    ENHET_FIELDS = [
        "navn", "organisasjonsform.kode", "naeringskode1.kode", "naeringskode1.beskrivelse",
        "antallAnsatte", "forretningsadresse.kommune", "forretningsadresse.postnummer",
        "forretningsadresse.poststed", "stiftelsesdato", "sisteInnsendteAarsregnskap",
        "konkurs", "underAvvikling", "erIKonsern",
    ]
    orgnr_set = set(r["orgnr"] for r in rows)
    enhet_lookup = {}
    with gzip.open(GZ_FILE, "rt", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            org = row.get("organisasjonsnummer")
            if org in orgnr_set:
                enhet_lookup[org] = {k: row.get(k, "") for k in ENHET_FIELDS}

    for r in rows:
        enhet = enhet_lookup.get(r["orgnr"], {})
        for k in ENHET_FIELDS:
            val = enhet.get(k, "")
            if k == "antallAnsatte" and val:
                try:
                    val = int(val)
                except ValueError:
                    pass
            r[k] = val

    enhet_headers = [
        "orgnr", "navn", "organisasjonsform.kode", "naeringskode1.kode", "naeringskode1.beskrivelse",
        "antallAnsatte", "forretningsadresse.kommune", "forretningsadresse.postnummer",
        "forretningsadresse.poststed", "stiftelsesdato", "sisteInnsendteAarsregnskap",
        "konkurs", "underAvvikling", "erIKonsern",
        "title", "panthaver", "panthaver_orgnr", "pantsetter", "pantsetter_orgnr",
        "beløp", "document_number", "innkomst_dato", "innkomst_kl", "rettsstiftelsen_er",
        "formuesgode", "krav", "total_pant_hos_bank", "total_beløp_hos_bank",
    ]
    local_xlsx = "/tmp/output_enhet.xlsx"
    n = write_xlsx(rows, enhet_headers, local_xlsx, fmt)
    gcs_xlsx = f"{GCS_PREFIX}/output/sparebanken_enhet.xlsx"
    gcs_upload(local_xlsx, gcs_xlsx)
    update_status("done", progress=f"{n} rows",
                  detail=f"gs://{BUCKET_NAME}/{gcs_xlsx}")


# ═══════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════

def main():
    """Pipeline entry point: download → collect → parse → export.

    Runs all 4 stages sequentially.  On error, writes error status
    to GCS before re-raising.
    """
    try:
        update_status("starting")
        download_enhetsregisteret()
        collect_all()
        parsed = parse_all()
        export_all(parsed)
        export(parsed)
    except Exception as e:
        update_status("error", error=str(e))
        raise


if __name__ == "__main__":
    main()
